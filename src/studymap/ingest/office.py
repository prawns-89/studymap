"""PPTX (one unit per slide, with speaker notes and pictures) and DOCX (one unit per heading section)."""
from __future__ import annotations

import hashlib
from collections import Counter
from pathlib import Path

from .common import char_count, clean_text, first_line, save_image_bytes
from .pdf import LOW_TEXT
from .records import Figure, Unit

DOCX_CHUNK = 4000      # characters per unit when a document has no headings


# --------------------------------------------------------------------------- pptx

def _shapes(shapes):
    """Every shape, descending into groups, in rough reading order (top, then left)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            out.extend(_shapes(sh.shapes))
        else:
            out.append(sh)
    return sorted(out, key=lambda s: ((s.top or 0) // 20000, s.left or 0))


def _table_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [[c.replace("|", "/").replace("\n", " ").strip() for c in r] + [""] * (w - len(r)) for r in rows]
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "|" + "---|" * w
    return "\n".join([head, sep] + ["| " + " | ".join(r) + " |" for r in rows[1:]])


def extract_pptx(path: Path, rel: str, kind: str, img_dir: Path, img_rel: str):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    warnings: list[str] = []
    slides = list(prs.slides)
    pics: list[list[tuple[str, bytes]]] = []
    seen = Counter()
    for slide in slides:
        lst = []
        for sh in _shapes(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob
                except Exception:
                    continue
                lst.append((hashlib.sha1(blob).hexdigest(), blob))
        pics.append(lst)
        seen.update({d for d, _ in lst})
    n = len(slides)
    deco = {d for d, c in seen.items() if n >= 3 and c >= n * 0.5}
    units = []
    for i, slide in enumerate(slides, 1):
        lines = []
        title_sh = slide.shapes.title
        title = clean_text(title_sh.text) if title_sh is not None and title_sh.has_text_frame else ""
        if title:
            lines.append(title)
        for sh in _shapes(slide.shapes):
            if title_sh is not None and sh.shape_id == title_sh.shape_id:
                continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = clean_text("".join(r.text for r in p.runs) or p.text)
                    if t:
                        lines.append("  " * (p.level or 0) + t)
            elif getattr(sh, "has_table", False) and sh.has_table:
                rows = [[clean_text(c.text) for c in r.cells] for r in sh.table.rows]
                lines.append(_table_md(rows))
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                name = clean_text(ch.chart_title.text_frame.text) if ch.has_title else ""
                lines.append(f"[chart{': ' + name if name else ''}]")
        text = "\n".join(lines).strip()
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
        chars = char_count(text)
        flags = ["low-text"] if chars < LOW_TEXT else []
        if slide._element.get("show") == "0":
            flags.append("hidden")
        figs, k = [], 0
        for d, blob in pics[i - 1]:
            if d in deco:
                continue
            k += 1
            name = f"s{i:03d}-{k}.jpg"
            wh = save_image_bytes(blob, img_dir / name)
            if wh is None:
                warnings.append(f"s{i}: picture {k} is in a format that can't be converted (WMF/EMF?)")
                continue
            figs.append(Figure(path=f"{img_rel}/{name}", why="embedded", w=wh[0], h=wh[1]))
        if figs:
            flags.append("figure")
        units.append(Unit(ref=f"{rel}#s{i}", n=i, title=first_line(title or text), text=text, chars=chars,
                          flags=flags, notes=notes, figures=figs))
    if any("low-text" in u.flags and not u.figures for u in units):
        warnings.append("some slides have little text and no extractable picture; "
                        "export the deck to PDF to get page renders")
    return units, {"slides": n}, warnings


# --------------------------------------------------------------------------- docx

def _heading_level(style_name: str) -> int:
    s = (style_name or "").lower()
    if s == "title":
        return 1
    if s.startswith("heading"):
        tail = s[7:].strip()
        return int(tail) if tail.isdigit() else 1
    return 0


def extract_docx(path: Path, rel: str, kind: str, img_dir: Path, img_rel: str):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    warnings: list[str] = []
    sections: list[dict] = []
    cur = {"title": "", "lines": [], "imgs": []}

    def flush():
        nonlocal cur
        if cur["lines"] or cur["imgs"] or cur["title"]:
            sections.append(cur)
        cur = {"title": "", "lines": [], "imgs": []}

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = clean_text(p.text)
            lvl = _heading_level(p.style.name if p.style is not None else "")
            for rid in child.xpath(".//a:blip/@r:embed"):
                part = d.part.related_parts.get(rid)
                if part is not None and hasattr(part, "blob"):
                    cur["imgs"].append(part.blob)
            if not text:
                continue
            if lvl:
                flush()
                cur["title"] = text
                cur["lines"].append("#" * min(lvl + 1, 6) + " " + text)
                continue
            is_list = "list" in (p.style.name.lower() if p.style is not None else "") or bool(child.xpath("./w:pPr/w:numPr"))
            cur["lines"].append(("- " if is_list else "") + text)
            if sum(len(x) for x in cur["lines"]) > DOCX_CHUNK and not cur["title"]:
                flush()
        elif tag == "tbl":
            t = Table(child, d)
            rows = []
            for r in t.rows:
                cells, last = [], None
                for c in r.cells:
                    if c._tc is not last:      # merged cells repeat; keep one copy
                        cells.append(clean_text(c.text))
                    last = c._tc
                rows.append(cells)
            cur["lines"].append(_table_md(rows))
    flush()
    units = []
    for i, s in enumerate(sections, 1):
        text = "\n\n".join(s["lines"]).strip()
        figs = []
        for k, blob in enumerate(s["imgs"], 1):
            name = f"sec{i:03d}-{k}.jpg"
            wh = save_image_bytes(blob, img_dir / name)
            if wh is None:
                warnings.append(f"sec{i}: image {k} can't be converted")
                continue
            figs.append(Figure(path=f"{img_rel}/{name}", why="embedded", w=wh[0], h=wh[1]))
        units.append(Unit(ref=f"{rel}#sec{i}", n=i, title=s["title"], text=text, chars=char_count(text),
                          flags=["figure"] if figs else [], figures=figs))
    return units, {"sections": len(units)}, warnings
