"""Shared fixtures: a synthetic course folder with one of everything ingest must handle."""
from __future__ import annotations

import io
import json
import shutil
from pathlib import Path

import pymupdf
import pytest

REPO = Path(__file__).resolve().parents[1]
REF = REPO / "studymap-starter" / "reference"
EXAMPLES = REF / "examples"

BULLETS = [
    "Round robin gives each process one quantum q, then moves it to the back of the ready queue.",
    "A small quantum means more context switches and more overhead for the scheduler.",
    "As the quantum grows without bound, round robin behaves exactly like FCFS scheduling.",
]


def png(rgb=(200, 80, 40), w=240, h=160) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), 0)
    pix.set_rect(pix.irect, rgb)
    pix.set_rect(pymupdf.IRect(10, 10, w // 2, h // 2), (rgb[2], rgb[0], rgb[1]))
    return pix.tobytes("png")


def transparent_png(w=300, h=200) -> bytes:
    """A diagram with a transparent background, like a draw.io export."""
    samples = bytearray(w * h * 4)                                 # all (0, 0, 0, 0): transparent
    for y in range(20, 120):
        for x in range(20, 120):
            samples[(y * w + x) * 4 + 3] = 255                     # one opaque black box
    return pymupdf.Pixmap(pymupdf.csRGB, w, h, bytes(samples), True).tobytes("png")


def make_deck(path: Path) -> None:
    """A 6-page slide PDF: template logo and footer on every page, a shadowed title, an incremental build,
    a vector figure, a big picture, and a page whose table numbers must survive footer removal."""
    doc = pymupdf.open()
    logo = png((10, 60, 160), 64, 64)
    pages = [
        dict(title="CPU Scheduling", body=["Instructor: A. Teacher"], shadow=True),
        dict(title="Round robin", body=BULLETS[:2]),
        dict(title="Round robin", body=BULLETS),                       # superset of page 2: page 2 is a build step
        dict(title="Process states", body=["new, ready, running, waiting, terminated"], drawings=True),
        dict(title="Gantt chart", body=[], picture=True),
        dict(title="Waiting time table", body=["P1 waits 0", "P2 waits 5", "P3 waits 12",
                                               "Average waiting time is (0 + 5 + 12) / 3 = 5.67 ms."]),
    ]
    for i, p in enumerate(pages, 1):
        page = doc.new_page(width=960, height=540)
        page.insert_image(pymupdf.Rect(880, 10, 940, 70), stream=logo)              # template decoration
        page.insert_text((40, 525), "CS F999: Testing Systems", fontsize=9)          # running footer
        page.insert_text((900, 525), str(i), fontsize=9)                              # page number
        if p.get("shadow"):
            page.insert_text((41, 81), p["title"], fontsize=30, color=(0.6, 0.6, 0.6))
        page.insert_text((40, 80), p["title"], fontsize=30)
        y = 130
        for b in p["body"]:
            page.insert_text((60, y), b, fontsize=14)
            y += 30
        if p.get("drawings"):
            for k in range(90):
                page.draw_line((300 + k * 4, 300), (300 + k * 4, 450))
        if p.get("picture"):
            page.insert_image(pymupdf.Rect(100, 120, 800, 480), stream=png((30, 160, 90), 600, 300))
    doc.save(path)


def make_text_pdf(path: Path, lines: list[str]) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for ln in lines:
        page.insert_text((72, y), ln, fontsize=11)
        y += 18
    doc.save(path)


def make_scanned_pdf(path: Path, pages: int = 2) -> None:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_image(page.rect, stream=png((235, 235, 225 - i * 20), 600, 840))
    doc.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    logo = png((10, 60, 160), 64, 64)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Paging {i + 1}" + ("\vVirtual memory, part 1" if i == 0 else "")
        body = s.placeholders[1].text_frame
        body.text = "A page table maps virtual page numbers to physical frame numbers."
        p = body.add_paragraph()
        p.text = "The TLB caches recent translations."
        p.level = 1
        s.shapes.add_picture(io.BytesIO(logo), Inches(9), Inches(0.1), Inches(0.5), Inches(0.5))
        if i == 1:
            t = s.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
            t.cell(0, 0).text, t.cell(0, 1).text = "Bits", "Meaning"
            t.cell(1, 0).text, t.cell(1, 1).text = "V", "valid"
            s.notes_slide.notes_text_frame.text = "Stress that the TLB is a cache."
        if i == 2:
            s.shapes.add_picture(io.BytesIO(png((120, 30, 160), 400, 300)), Inches(1), Inches(3), Inches(4), Inches(3))
    prs.save(str(path))


def make_docx(path: Path) -> None:
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.add_heading("Deadlock", 1)
    d.add_paragraph("Four conditions must hold at once for a deadlock to occur.")
    d.add_paragraph("Mutual exclusion", style="List Bullet")
    d.add_paragraph("Hold and wait", style="List Bullet")
    d.add_heading("Banker's algorithm", 1)
    d.add_paragraph("It grants a request only if the resulting state is safe.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Process", "Need"
    t.cell(1, 0).text, t.cell(1, 1).text = "P0", "7 4 3"
    d.add_picture(io.BytesIO(png((90, 90, 200), 300, 200)), width=Inches(2))
    d.save(str(path))


MIPS = """# sum 1..10
        .text
main:   li   $t0, 0
        li   $t1, 1
loop:   add  $t0, $t0, $t1
        addi $t1, $t1, 1
        ble  $t1, 10, loop
        li   $v0, 10
        syscall
"""
X86 = """.code16
start:  movw %ax, %ds
        movl %eax, %ebx
        int $0x13
"""
C_SRC = '#include <stdio.h>\n\nint main(void) {\n    printf("hi\\n");\n    return 0;\n}\n'


def build_course(root: Path) -> Path:
    for d in ("slides", "notes", "labs/lab1/build", "papers", "extra", "tutorials", ".git"):
        (root / d).mkdir(parents=True, exist_ok=True)
    (root / "course.yaml").write_text(
        "name: Testing Systems\ncode: CS F999\nexams:\n  - name: midsem\n    weight: 30\n"
        "ignore: [\"labs/lab1/build\"]\n", encoding="utf-8")
    make_deck(root / "slides" / "07-sched.pdf")
    make_pptx(root / "slides" / "08-paging.pptx")
    make_docx(root / "notes" / "deadlock.docx")
    (root / "notes" / "summary.md").write_text(
        "# Scheduling\n\nFCFS, SJF, RR.\n\n```python\n# not a heading\nprint(1)\n```\n\n## Paging\n\nPages and frames.\n",
        encoding="utf-8")
    (root / "tutorials" / "t1.txt").write_text("Tutorial 1: compute the average waiting time.\n", encoding="utf-8")
    (root / "labs" / "lab1" / "main.c").write_text(C_SRC, encoding="utf-8")
    (root / "labs" / "lab1" / "sum.asm").write_text(MIPS, encoding="utf-8")
    (root / "labs" / "lab1" / "boot.S").write_text(X86, encoding="utf-8")
    (root / "labs" / "lab1" / "Makefile").write_text("all:\n\tgcc -o main main.c\n", encoding="utf-8")
    (root / "labs" / "lab1" / "main.o").write_bytes(b"\x7fELF\x00\x01")
    (root / "labs" / "lab1" / "build" / "gen.c").write_text("int x;\n", encoding="utf-8")
    (root / "labs" / "lab1" / "diagram.png").write_bytes(png((50, 50, 50), 300, 300))
    (root / "labs" / "lab1" / "datapath.png").write_bytes(transparent_png())
    make_scanned_pdf(root / "papers" / "2024-midsem.pdf")
    make_text_pdf(root / "papers" / "2024-midsem-key.pdf", ["Q1. Average waiting time = 5.67 ms [3]", "56", "156"])
    make_text_pdf(root / "papers" / "2023-compre-sol.pdf", ["Solutions without a paper"])
    make_text_pdf(root / "papers" / "sample-questions.pdf", ["Q1. Explain round robin. [4]"])
    (root / "extra" / "old.zip").write_bytes(b"PK\x03\x04")
    (root / "extra" / "deck.ppt").write_bytes(b"\xd0\xcf\x11\xe0")
    (root / ".git" / "config").write_text("[core]\n", encoding="utf-8")
    return root


@pytest.fixture(scope="session")
def course_template(tmp_path_factory) -> Path:
    """Built once per session; tests copy it so they can change files freely."""
    return build_course(tmp_path_factory.mktemp("template") / "CS-F999")


@pytest.fixture
def course(course_template, tmp_path) -> Path:
    dst = tmp_path / "CS-F999"
    shutil.copytree(course_template, dst)
    return dst


# --------------------------------------------------------------------------- a small paper analysis

TOPICS = {
    "clusters": [{"id": "sched", "name": "Scheduling"}, {"id": "mem", "name": "Memory"}],
    "topics": [
        {"id": "rr", "name": "Round robin", "cluster": "sched", "sources": ["slides/07-sched.pdf#p2-3"]},
        {"id": "states", "name": "Process states", "cluster": "sched", "sources": ["slides/07-sched.pdf#p4"]},
        {"id": "metrics", "name": "Scheduling metrics", "cluster": "sched", "sources": ["slides/07-sched.pdf#p5-6"]},
        {"id": "paging", "name": "Paging", "cluster": "mem", "sources": ["slides/08-paging.pptx#s1-3"]},
        {"id": "ipc", "name": "IPC", "cluster": "mem", "sources": []},
    ],
    "uncovered_ok": [{"ref": "slides/07-sched.pdf#p1", "reason": "title slide"}],
}
PAPERS = {
    "papers": [
        {"id": "2024-midsem", "path": "papers/2024-midsem.pdf", "exam": "midsem", "year": 2024, "total_marks": 10,
         "keys": [{"path": "papers/2024-midsem-key.pdf"}], "transcribed_from_image": True},
        {"id": "samples", "path": "papers/sample-questions.pdf", "kind": "sample", "exam": "midsem"},
    ],
    "questions": [
        {"id": "2024-midsem-q1", "paper": "2024-midsem", "number": "1", "marks": 4, "type": "numerical", "mode": "apply",
         "difficulty": "medium", "topics": ["rr", "metrics"], "text": "Run RR with q = 2.", "src": "papers/2024-midsem.pdf#p1",
         "answer": {"text": "5.67 ms", "src": ["papers/2024-midsem-key.pdf#p1"]}},
        {"id": "2024-midsem-q2", "paper": "2024-midsem", "number": "2", "marks": 6, "type": "long-explain",
         "mode": "understand", "difficulty": "easy", "topics": ["ipc"], "text": "Compare pipes and shared memory.",
         "src": "papers/2024-midsem.pdf#p2"},
        {"id": "samples-1", "paper": "samples", "number": "1", "type": "short", "mode": "understand", "difficulty": "easy",
         "topics": ["paging"], "text": "Explain round robin.", "src": "papers/sample-questions.pdf#p1"},
    ],
    "patterns": [{"text": "Every paper has a scheduling numerical.", "questions": ["2024-midsem-q1"]}],
}


def write_analysis(course, topics=TOPICS, papers=PAPERS):
    d = course / "_studymap" / "analysis"
    d.mkdir(parents=True, exist_ok=True)
    (d / "topics.json").write_text(json.dumps(topics, indent=1))
    if papers is not None:
        (d / "papers.json").write_text(json.dumps(papers, indent=1))
    return d


@pytest.fixture
def ingested_course(course):
    from studymap.ingest import ingest
    ingest(course, log=lambda s: None)
    return course
